from fastapi import APIRouter, HTTPException
from pydantic import BaseModel
from typing import List, Dict, Any, Optional
import logging
import re
import unicodedata
from pathlib import Path

from services.extract_service import extract_info_via_ai
from services.history_service import record_extraction
from services.scheme_service import touch_scheme
from services.upload_service import UPLOAD_META, parse_file, parse_file_with_trace
from models.schemas import ExtractRequest, ExtractResponseItem
from db.database import SessionLocal
from db.models import Asset

router = APIRouter(prefix="/api", tags=["extract"])

logger = logging.getLogger(__name__)


def _to_str(value: Any) -> str:
    """把任意值安全地转为 UTF-8 可解码的字符串。

    关键：若上游误把文件二进制（bytes）传进来，绝不能让它直接进入异常消息或
    响应字段——否则 FastAPI 在 JSON 序列化时会对 bytes 执行 utf-8 解码，触发
    UnicodeDecodeError（byte 0xb2 ...）导致 500。这里统一用 errors='replace'
    清洗，保证后续任何拼接/返回都是纯文本。"""
    if isinstance(value, bytes):
        try:
            return value.decode("utf-8", errors="replace")
        except Exception:
            return "<binary data>"
    if value is None:
        return ""
    return str(value)


def _sanitize_documents(documents: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
    """入口清洗：确保每个文档的 filename / content 等都是干净字符串，
    从源头杜绝二进制流入 LLM 提示词与最终响应。"""
    cleaned: List[Dict[str, Any]] = []
    for doc in documents or []:
        if not isinstance(doc, dict):
            continue
        new_doc: Dict[str, Any] = {}
        for k, v in doc.items():
            # 仅把这些字段当文本清洗；其余字段原样保留（但仍避免 bytes）
            if k in ("filename", "content", "file_type", "role"):
                new_doc[k] = _to_str(v)
            elif isinstance(v, bytes):
                new_doc[k] = _to_str(v)
            else:
                new_doc[k] = v
        cleaned.append(new_doc)
    return cleaned


def _asset_path_by_name(fn: str) -> Optional[str]:
    """UPLOAD_META 为内存登记，后端重启即丢失；回退到资产库按原始文件名
    查磁盘路径（取最新一条且文件仍存在），保证重启后溯源仍可用。"""
    if not fn:
        return None
    try:
        with SessionLocal() as session:
            asset = (
                session.query(Asset)
                .filter(Asset.original_name == fn)
                .order_by(Asset.id.desc())
                .first()
            )
            if asset and Path(asset.path).exists():
                return asset.path
    except Exception as exc:
        logger.warning("资产库回退查询 %s 失败: %s", fn, exc)
    return None


def _build_trace_map(
    documents: List[Dict[str, Any]]
) -> Dict[str, Dict[str, Any]]:
    """为每个文档生成溯源段落：按文件名找到已上传的源文件，用 parse_file_with_trace
    重新解析。溯源与上传时的纯文本来自同一份文件，保证来源一致、偏移自洽。

    任务4：返回结构为 {文件名: {"chunks": [段落...], "para_ids": {}}}。
    chunks 是 ParseChunk.to_dict() 列表；para_ids 为预留字段（LLM 返回的段落ID
    映射，当前按单元格经 ||p:N|| 携带，故此处留空字典占位，保持结构约定）。"""
    by_name: Dict[str, Dict[str, Any]] = {}
    for meta in UPLOAD_META.values():
        if meta.get("original_name"):
            by_name[meta["original_name"]] = meta  # 后者覆盖 = 取最新
    trace: Dict[str, Dict[str, Any]] = {}
    for doc in documents:
        fn = doc.get("filename")
        meta = by_name.get(fn)
        path = meta.get("path") if meta else None
        if not path:
            path = _asset_path_by_name(fn)
            if path:
                logger.info("文档 %s 不在 UPLOAD_META，回退资产库路径 %s", fn, path)
        if not path:
            logger.warning("文档 %s 未在 UPLOAD_META 中登记源文件路径，溯源将被跳过", fn)
            continue
        try:
            chunks = parse_file_with_trace(path)
            trace[fn] = {
                "chunks": [c.to_dict() for c in chunks],
                "para_ids": {},
            }
        except Exception as e:
            logger.warning("为文档 %s 生成溯源失败: %s", fn, e)
    return trace


@router.post("/extract", response_model=List[ExtractResponseItem])
async def extract_info(req: ExtractRequest):
    # 入口清洗：把 filename/content 等强制转为干净字符串，杜绝文件二进制
    # 流入 LLM 提示词或最终响应（否则 JSON 序列化时会因非 UTF-8 字节崩溃成 500）。
    documents = _sanitize_documents(req.documents)
    try:
        trace = _build_trace_map(documents)
        results = await extract_info_via_ai(req.prompt, req.fields, documents, trace=trace)
    except Exception as e:
        # 只记录“文件名 + 错误类型”等元信息，绝不把文件内容/二进制拼进异常消息，
        # 保证抛出的 detail 永远是可 JSON 序列化的纯文本。
        names = [d.get("filename", "<unknown>") for d in documents]
        err_type = type(e).__name__
        logger.error("提取失败 files=%s error=%s", names, err_type)
        # 注意：detail 只用错误类型，不使用 str(e)，避免其内含二进制再次触发解码错误
        raise HTTPException(
            status_code=500,
            detail=f"提取失败({err_type})，涉及文件：{', '.join(names)}",
        ) from e

    # 提取完成即落一条任务历史（配置与结果存快照）；失败只记日志，不影响响应
    template_meta: Optional[Dict[str, Any]] = None
    for meta in UPLOAD_META.values():
        if meta.get("role") == "template":
            template_meta = meta  # 遍历取最后出现的模板记录
    task_id = record_extraction(
        prompt=req.prompt,
        fields=req.fields,
        results=[item.model_dump() for item in results],
        source_filenames=[str(d.get("filename", "")) for d in documents],
        template_asset_id=(template_meta or {}).get("asset_id"),
        template_name=(template_meta or {}).get("original_name", ""),
    )
    if task_id is not None:
        for item in results:
            item.task_id = str(task_id)

    # 回写方案使用时间（用于推荐接口的“最近使用”排序）；失败只记日志
    if req.scheme_id and str(req.scheme_id).isdigit():
        touch_scheme(int(req.scheme_id))
    return results


# ========== 智能提取模板表头接口 ==========
@router.get("/extract/template-headers")
async def get_template_headers():
    """
    获取最新上传的模板文件的表头信息。
    支持带标题/说明文字的复杂 Excel，自动定位真实表头行。
    如果未找到模板或解析失败，返回空列表。
    """
    # 1. 找到最新的模板文件（遍历取最后出现的 role=template 记录）
    latest_template_meta: Optional[Dict[str, Any]] = None
    for meta in UPLOAD_META.values():
        if meta.get("role") == "template":
            latest_template_meta = meta

    if not latest_template_meta:
        logger.warning("未找到 role=template 的上传记录")
        return {"headers": []}

    logger.info("选中的模板记录: %s", latest_template_meta)

    # 2. 解析模板文件内容
    try:
        path = Path(latest_template_meta["path"])
        if not path.exists():
            logger.warning("模板文件不存在: %s", path)
            return {"headers": []}

        # 统一处理扩展名：兼容 "xlsx" / ".xlsx" / "XLSX" 等写法
        file_ext = str(latest_template_meta.get("ext", "")).lower().strip()
        if not file_ext.startswith('.'):
            file_ext = '.' + file_ext

        headers: List[str] = []

        if file_ext in ('.xlsx', '.xlsm'):
            headers = _extract_xlsx_headers(path)
        elif file_ext == '.docx':
            headers = _extract_docx_headers(path)
        else:
            # 非 Excel / Word（CSV / 文本等）走通用文本解析
            status, content, error = parse_file(path, file_ext)
            if status != "success":
                logger.warning("模板解析失败: %s", error)
                return {"headers": []}

            lines = content.strip().split('\n')
            if not lines:
                logger.info("内容为空，无法提取表头")
                return {"headers": []}

            headers = re.split(r'[,\t|;]+', lines[0].strip())
            headers = [h.strip() for h in headers if h.strip()]

        if not headers:
            logger.info("已读取文件，但未解析出任何有效表头")
            return {"headers": []}

        logger.info("成功提取到模板表头: %s", headers)
        return {"headers": headers}

    except HTTPException:
        raise
    except Exception as e:
        logger.exception("获取模板表头时发生未知错误: %s", e)
        raise HTTPException(status_code=500, detail=f"服务器内部错误: {str(e)}")


# ========== 字段自动匹配接口 ==========
class AutoMatchRequest(BaseModel):
    template_fields: List[str]
    extract_fields: List[str]


def _norm_field(s: Any) -> str:
    """字段名归一化：全角转半角、去全部空白、转小写。
    使 'GDP总量（亿元）' 与 'GDP总量(亿元) ' 等写法差异不影响匹配。"""
    s = unicodedata.normalize("NFKC", _to_str(s))
    return re.sub(r"\s+", "", s).lower()


@router.post("/extract/auto-match-fields")
async def auto_match_fields(req: AutoMatchRequest):
    """模板字段 → 提取字段自动匹配：归一化后精确同名者一一对应。

    规则：
    - 仅做确定性精确匹配（归一化后字符串相等），不做模糊/语义猜测；
    - 一一对应：每个提取字段最多被占用一次，避免多对一错配；
    - 匹配不上的模板字段不出现在结果中，由用户手动选择。
    """
    tmpl_fields = [f for f in (req.template_fields or []) if _to_str(f).strip()]
    ext_fields = [f for f in (req.extract_fields or []) if _to_str(f).strip()]

    # 提取字段按归一化名建索引（保序，重名取首个）
    ext_index: Dict[str, str] = {}
    for f in ext_fields:
        key = _norm_field(f)
        if key and key not in ext_index:
            ext_index[key] = f

    mapping: Dict[str, str] = {}
    used = set()
    for tf in tmpl_fields:
        if tf in mapping:
            continue  # 模板内重名字段只映射一次
        key = _norm_field(tf)
        hit = ext_index.get(key)
        if hit is not None and hit not in used:
            mapping[tf] = hit
            used.add(hit)

    logger.info(
        "字段自动匹配: 模板 %d 个, 提取 %d 个, 命中 %d 个",
        len(tmpl_fields), len(ext_fields), len(mapping),
    )
    return {
        "mapping": mapping,
        "matched": len(mapping),
        "total": len(dict.fromkeys(tmpl_fields)),
    }


def _extract_xlsx_headers(path: Path) -> List[str]:
    """
    Excel 表头提取：自动定位真实表头行（非空单元格最多且达到阈值的一行），
    并用左侧值填充合并单元格造成的空值。
    说明：刻意使用项目已有依赖 openpyxl，避免为单个接口引入 pandas。
    """
    # 延迟导入：可选格式依赖禁止模块顶层 import
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=False, data_only=True)
    try:
        ws = wb[wb.sheetnames[0]]
        # 先把合并单元格的值回填到每个被合并位置：
        # 否则表头行因合并出现空洞，"非空最多"启发式会错选下方的满列数据行
        grid = [list(row) for row in ws.iter_rows(values_only=True)]
        for rng in ws.merged_cells.ranges:
            top_left = grid[rng.min_row - 1][rng.min_col - 1]
            for r in range(rng.min_row - 1, rng.max_row):
                for c in range(rng.min_col - 1, rng.max_col):
                    grid[r][c] = top_left
        rows = [tuple(row) for row in grid]
    finally:
        wb.close()

    if not rows:
        return []

    logger.info("Excel 读取成功，共 %d 行", len(rows))

    # 智能定位真实表头行：非空单元格数量最多且达到阈值的一行
    min_threshold = 3
    best_idx = -1
    max_cols_count = 0
    for i, row in enumerate(rows):
        non_empty = sum(1 for cell in row if cell is not None and str(cell).strip())
        if non_empty > max_cols_count and non_empty >= min_threshold:
            max_cols_count = non_empty
            best_idx = i

    if best_idx == -1:
        logger.info("未找到满足阈值(%d列)的表头行", min_threshold)
        return []

    logger.info("定位到表头行索引: %d，非空列数: %d", best_idx, max_cols_count)

    # 用左侧值填充合并单元格造成的空值（ffill），并清理换行/首尾空格
    headers: List[str] = []
    last_val = ""
    for cell in rows[best_idx]:
        if cell is not None and str(cell).strip():
            last_val = str(cell).replace('\n', '').replace('\r', '').strip()
        if last_val:
            headers.append(last_val)
    return headers


def _extract_docx_headers(path: Path) -> List[str]:
    """Word 表头提取：遍历所有表格，取列数最多的主表首行作为表头。"""
    # 延迟导入：可选格式依赖禁止模块顶层 import
    from docx import Document

    doc = Document(str(path))
    if not doc.tables:
        logger.info("Word 文档中未发现任何表格")
        return []

    best_table = None
    max_cols = 0
    for table in doc.tables:
        if not table.rows:
            continue
        cols = len(table.rows[0].cells)
        if cols > max_cols:
            max_cols = cols
            best_table = table

    if not best_table:
        logger.info("Word 文档中没有有效的表格结构")
        return []

    headers: List[str] = []
    for cell in best_table.rows[0].cells:
        text = cell.text.replace('\n', '').replace('\r', '').strip()
        if text:
            headers.append(text)

    if not headers:
        logger.info("Word 表格第一行未解析到任何文字")
        return []

    logger.info("从 Word 表格成功提取表头: %s", headers)
    return headers
