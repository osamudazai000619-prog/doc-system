import logging
import re
import uuid
import zipfile
from collections import Counter
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple
from xml.etree import ElementTree as ET

from fastapi import UploadFile
from starlette.concurrency import run_in_threadpool

from docx import Document
from openpyxl import load_workbook

from services.asset_service import register_asset

BASE_DIR = Path("uploads")
TARGET_DIR = BASE_DIR / "targets"
TEMPLATE_DIR = BASE_DIR / "templates"
for _d in (TARGET_DIR, TEMPLATE_DIR):
    _d.mkdir(parents=True, exist_ok=True)

ALLOWED_EXT = {".docx", ".txt", ".md", ".xlsx", ".pdf"}
MAX_FILE_SIZE = 20 * 1024 * 1024  # 20MB

# 已上传文件的元数据（file_id -> meta）。export 服务依赖它按模板名查找文件路径。
# 单进程内存态即可满足当前 demo；生产环境建议替换为 Redis / 数据库
UPLOAD_META: Dict[str, Dict[str, Any]] = {}


def safe_filename(name: str) -> str:
    """清理文件名，防止路径穿越和奇怪字符。"""
    name = Path(name or "unnamed").name
    return re.sub(r"[^\w\u4e00-\u9fa5.\-]+", "_", name)


def save_upload_file(upload: UploadFile, role: str) -> Tuple[str, Path, str, str]:
    """
    保存上传文件到对应目录。
    role: target / template
    返回: file_id, save_path, original_name, ext
    """
    original = safe_filename(upload.filename or "unnamed")
    ext = Path(original).suffix.lower()
    if ext not in ALLOWED_EXT:
        raise ValueError(f"不支持的文件格式：{ext or '未知'}，仅支持 DOCX/TXT/MD/XLSX/PDF")

    file_id = uuid.uuid4().hex
    save_dir = TARGET_DIR if role == "target" else TEMPLATE_DIR
    save_path = save_dir / f"{file_id}{ext}"

    size = 0
    with save_path.open("wb") as f:
        while True:
            chunk = upload.file.read(1024 * 1024)
            if not chunk:
                break
            size += len(chunk)
            if size > MAX_FILE_SIZE:
                f.close()
                save_path.unlink(missing_ok=True)
                raise ValueError(f"文件超过大小限制（>{MAX_FILE_SIZE // 1024 // 1024}MB）")
            f.write(chunk)

    asset_id, effective_path = register_asset(save_path, original, ext, role, size)
    UPLOAD_META[file_id] = {
        "file_id": file_id,
        "path": str(effective_path),
        "original_name": original,
        "role": role,
        "ext": ext,
        "size": size,
        "asset_id": asset_id,
    }
    return file_id, effective_path, original, ext


def read_text_file(path: Path) -> str:
    """读取 txt/md，尝试多种编码。"""
    raw = path.read_bytes()
    for enc in ("utf-8", "utf-8-sig", "gb18030", "gbk", "big5", "latin-1"):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="replace")


def _parse_docx_raw(path: Path) -> str:
    """
    兜底解析：docx 因缺失零件（如 footnotes.xml）导致 python-docx 打不开时，
    直接解包读取主文档 XML，提取全部段落文本（含表格内段落）。
    """
    ns = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    with zipfile.ZipFile(path) as z:
        xml = z.read("word/document.xml")
    root = ET.fromstring(xml)
    parts = []
    for para in root.iter(f"{ns}p"):
        text = "".join(t.text or "" for t in para.iter(f"{ns}t")).strip()
        if text:
            parts.append(text)
    return "\n".join(parts)


# ==================== PDF 富文本解析 ====================
# 相比裸 page.get_text()，重点解决四类边界问题（策略见各函数注释）：
#   1. 块级丢失：文本永不因被背景色/边框图形覆盖而丢弃；提示框（实心矩形背景）
#      内的文字按坐标归属输出为独立 blockquote 段落
#   2. 上下标压平：span 上标 flag + 「字号缩小 + 基线偏移」启发式，还原 Unicode 上下标
#   3. 代码块空白丢失：等宽特征或「描边容器 + 代码特征」命中时按行原样保留
#   4. 列表符号丢失：文本 bullet 归一化 + 矢量小圆点坐标对齐还原

# Unicode 上/下标映射源表：仅收录存在标准映射的字符；含未收录字符时退化为 ^x / _x
_PDF_SUP_SRC: Dict[str, str] = {
    "0": "\u2070", "1": "\u00b9", "2": "\u00b2", "3": "\u00b3", "4": "\u2074",
    "5": "\u2075", "6": "\u2076", "7": "\u2077", "8": "\u2078", "9": "\u2079",
    "+": "\u207a", "-": "\u207b", "=": "\u207c", "(": "\u207d", ")": "\u207e",
    "a": "\u1d43", "b": "\u1d47", "c": "\u1d9c", "d": "\u1d48", "e": "\u1d49",
    "f": "\u1da0", "g": "\u1d4d", "h": "\u02b0", "i": "\u2071", "j": "\u02b2",
    "k": "\u1d4f", "l": "\u02e1", "m": "\u1d50", "n": "\u207f", "o": "\u1d52",
    "p": "\u1d56", "r": "\u02b3", "s": "\u02e2", "t": "\u1d57", "u": "\u1d58",
    "v": "\u1d5b", "w": "\u02b7", "x": "\u02e3", "y": "\u02b8", "z": "\u1dbb",
}
_PDF_SUB_SRC: Dict[str, str] = {
    "0": "\u2080", "1": "\u2081", "2": "\u2082", "3": "\u2083", "4": "\u2084",
    "5": "\u2085", "6": "\u2086", "7": "\u2087", "8": "\u2088", "9": "\u2089",
    "+": "\u208a", "-": "\u208b", "=": "\u208c", "(": "\u208d", ")": "\u208e",
    "a": "\u2090", "e": "\u2091", "h": "\u2095", "i": "\u1d62", "j": "\u2c7c",
    "k": "\u2096", "l": "\u2097", "m": "\u2098", "n": "\u2099", "o": "\u2092",
    "p": "\u209a", "r": "\u1d63", "s": "\u209b", "t": "\u209c", "u": "\u1d64",
    "v": "\u1d65", "x": "\u2093",
}
_PDF_SUP_MAP = str.maketrans(_PDF_SUP_SRC)
_PDF_SUB_MAP = str.maketrans(_PDF_SUB_SRC)

# 等宽字体名特征：Skia/Word 等生成器产出的 PDF 可能既无 MONO flag 也无规范字体名，
# 此处仅作辅助信号（主判定见 _pdf_group_is_code）
_PDF_MONO_FONT_RE = re.compile(
    r"mono|courier|consol|menlo|cascadia|jetbrains|source[-._ ]?code|inconsolata|"
    r"fira[-._ ]?(code|mono)|dejavu[-._ ]?sans[-._ ]?mono|sarasa|pt[-._ ]?mono|"
    r"ubuntu[-._ ]?mono|roboto[-._ ]?mono|lucida[-._ ]?console|andale|hack\b", re.I)

# 代码行特征：常见关键字 / 行尾大括号分号 / 运算符 / 函数调用形态
_PDF_CODE_LINE_RE = re.compile(
    r"\b(?:def|class|function|return|import|export|from|if|elif|else|for|while|"
    r"try|except|catch|finally|const|let|var|print|printf|println|echo|"
    r"public|private|protected|static|void|struct|enum|package|fn|impl|use|"
    r"match|lambda|yield|await|async|typeof|instanceof)\b"
    r"|[{};]\s*$|=>|::|&&|\|\||==|!=|<=|>=|\w+\s*\([^)]*\)")

_PDF_BULLET_STRONG = "•●◦‣⁃▪▸■□◆◉∙○"
_PDF_BULLET_WEAK = "·・"  # 弱符号仅当后随空格才视为 bullet（避免误伤中文人名/章节号）
_PDF_CJK_RE = re.compile(r"[\u2e80-\u9fff\uf900-\ufaff\u3000-\u303f\uff00-\uffef\u3040-\u30ff\uac00-\ud7af]")


def _pdf_median(values: List[float]) -> float:
    vals = sorted(values)
    n = len(vals)
    if not n:
        return 0.0
    return vals[n // 2] if n % 2 else (vals[n // 2 - 1] + vals[n // 2]) / 2.0


def _pdf_line_text(line: dict) -> str:
    """
    行内文本还原（痛点 2：上下标压平）。
    判定：PyMuPDF span 的 bit0 flag 标记上标；下标无 flag，需按「字号 < 行主字号
    92% 且基线偏移超 10% 主字号」启发式补判（上标同理补判，覆盖无 flag 的情况）。
    主字号/主基线取行内按字符数加权的中位数，避免个别大字带偏基准。
    """
    spans = [s for s in line.get("spans", []) if s.get("text")]
    if not spans:
        return ""
    size_pop: List[float] = []
    base_pop: List[float] = []
    for s in spans:
        w = max(len(s["text"].strip()), 1)
        size_pop.extend([s.get("size", 11.0)] * w)
        origin = s.get("origin")
        base_pop.extend([(origin[1] if origin else s["bbox"][3])] * w)
    med_size = _pdf_median(size_pop) or 11.0
    med_base = _pdf_median(base_pop)
    out = []
    for s in spans:
        t = s["text"].replace("\u00a0", " ")
        size = s.get("size", 11.0)
        origin = s.get("origin")
        base = origin[1] if origin else s["bbox"][3]
        small = size < med_size * 0.92
        off = base - med_base
        sup_flag = bool(s.get("flags", 0) & 1)
        is_sup = sup_flag or (small and off < -0.10 * med_size)
        is_sub = (not sup_flag) and small and off > 0.10 * med_size
        if (is_sup or is_sub) and t.strip():
            table = _PDF_SUP_SRC if is_sup else _PDF_SUB_SRC
            if all(ch in table or ch.isspace() for ch in t):
                out.append(t.translate(_PDF_SUP_MAP if is_sup else _PDF_SUB_MAP))
            else:
                out.append(f"^{t.strip()}" if is_sup else f"_{t.strip()}")
        else:
            out.append(t)
    return "".join(out)


def _pdf_analyze_drawings(
    drawings: List[dict], page_w: float, page_h: float
) -> Tuple[List[Tuple[float, float, float, float]],
           List[Tuple[float, float, float, float]],
           List[Tuple[float, float]]]:
    """
    分类页面绘图对象（痛点 1、4 的图形侧）：
    - bg_rects：实心矩形填充（含 're' 项）且非页面底色/细线/超大 → 提示框背景候选。
      只认 're' 项是关键：整节内容的「边框环」在 PDF 中通常渲染为仅含 l/c 线段的
      环形路径（bbox 覆盖全节），若按 bbox 当背景会把整节文字误标为提示框。
    - stroke_rects：带描边的较大矩形 → 代码块容器候选。
    - dots：小尺寸近方形的实心路径 → 矢量无序列表符号候选。
    """
    page_area = page_w * page_h
    bg_rects: List[Tuple[float, float, float, float]] = []
    stroke_rects: List[Tuple[float, float, float, float]] = []
    dots: List[Tuple[float, float]] = []
    seen_bg = set()
    seen_dot = set()
    for dr in drawings:
        r = dr.get("rect")
        if r is None:
            continue
        w, h = r.width, r.height
        typ = dr.get("type", "")
        fill = dr.get("fill")
        items = dr.get("items") or []
        if typ in ("f", "fs") and fill is not None:
            solid_rect = any(it[0] == "re" for it in items)
            if (solid_rect
                    and not (w >= page_w * 0.95 and h >= page_h * 0.95)  # 页面底色
                    and min(w, h) > 2.5 and (max(w, h) / min(w, h)) <= 15  # 细线/装饰条
                    and w * h <= page_area * 0.15):  # 大面积底色不当提示框
                key = (round(r.x0), round(r.y0), round(r.x1), round(r.y1))
                if key not in seen_bg:
                    seen_bg.add(key)
                    bg_rects.append((r.x0, r.y0, r.x1, r.y1))
        if typ in ("s", "fs") and w > 40 and h > 30:
            stroke_rects.append((r.x0, r.y0, r.x1, r.y1))
        if (typ == "f" and fill is not None
                and 1.0 <= w <= 8.0 and 1.0 <= h <= 8.0 and abs(w - h) <= 2.0):
            ckey = (round((r.x0 + r.x1) / 4.0), round((r.y0 + r.y1) / 4.0))
            if ckey not in seen_dot:
                seen_dot.add(ckey)
                dots.append(((r.x0 + r.x1) / 2.0, (r.y0 + r.y1) / 2.0))
    return bg_rects, stroke_rects, dots


def _pdf_collect_blocks(pdict: dict, page_h: float, drop_zone: set, drop_rot: set) -> List[dict]:
    """
    收集文本块（痛点 1：只按「跨页重复」剔除页眉页脚与旋转水印，绝不按图形
    覆盖关系剔除，避免误伤提示框内的正常文本）。同时统计等宽字符占比供代码块判定。
    """
    blocks: List[dict] = []
    for b in pdict.get("blocks", []):
        if b.get("type") != 0:
            continue
        lines = []
        size_max = 0.0
        mono_chars = 0
        total_chars = 0
        for line in b.get("lines", []):
            text = _pdf_line_text(line)
            norm = re.sub(r"\s+", "", text)
            if not norm:
                continue
            bbox = line.get("bbox") or (0, 0, 0, 0)
            y0, y1 = bbox[1], bbox[3]
            dx, dy = line.get("dir", (1.0, 0.0))[:2]
            if abs(dy) <= 0.05:
                # 仅顶部/底部区域且跨页重复 → 页眉页脚
                if (y1 <= page_h * 0.07 or y0 >= page_h * 0.93) and norm in drop_zone:
                    continue
            elif norm in drop_rot:
                continue  # 旋转且重复出现 → 水印/印章
            lines.append({"text": text, "x0": bbox[0], "y0": y0, "x1": bbox[2], "y1": y1})
            for s in line.get("spans", []):
                t = s.get("text", "")
                if not t.strip():
                    continue
                size_max = max(size_max, s.get("size", 0.0))
                n_span = len(t.strip())
                total_chars += n_span
                if (s.get("flags", 0) & 8) or _PDF_MONO_FONT_RE.search(s.get("font", "")):
                    mono_chars += n_span
        if not lines:
            continue
        blocks.append({
            "x0": min(l["x0"] for l in lines), "y0": min(l["y0"] for l in lines),
            "x1": max(l["x1"] for l in lines), "y1": max(l["y1"] for l in lines),
            "size": size_max or 11.0, "lines": lines,
            "mono_chars": mono_chars, "total_chars": total_chars,
        })
    return blocks


def _pdf_group_is_code(group: List[dict]) -> bool:
    """容器内成组代码判定：等宽字符过半，或 ASCII 占比达标且代码特征行占比 ≥ 60%。"""
    nonblank = [l for b in group for l in b["lines"] if l["text"].strip()]
    if not nonblank:
        return False
    mono = sum(b["mono_chars"] for b in group)
    total = sum(b["total_chars"] for b in group)
    if total and mono / total >= 0.5:
        return True
    text_all = "\n".join(l["text"] for l in nonblank)
    if sum(1 for ch in text_all if ord(ch) < 128) / max(len(text_all), 1) < 0.4:
        return False
    hits = sum(1 for l in nonblank if _PDF_CODE_LINE_RE.search(l["text"]))
    return hits / len(nonblank) >= 0.6


def _pdf_render_code(group: List[dict]) -> str:
    """
    渲染围栏代码块（痛点 3）：行序按 y→x 排列；文本自带行首空格时原样保留，
    否则按与容器左缘的距离合成缩进；全程不做空白规整、不合并行。
    """
    lines = [l for b in group for l in b["lines"]]
    lines.sort(key=lambda l: (round(l["y0"], 1), l["x0"]))
    base_x = min(l["x0"] for l in lines)
    sizes = [s.get("size", 11.0) for b in group for l in b["lines"]
             for s in (l.get("spans") or []) if s.get("text", "").strip()]
    char_w = max((_pdf_median(sizes) if sizes else 10.0) * 0.6, 1.0)
    raw_lines: List[str] = []
    for l in lines:
        t = l["text"].rstrip()
        if not t.strip():
            raw_lines.append("")
            continue
        if t[:1] in (" ", "\t"):
            raw_lines.append(t.rstrip())  # 自带缩进：原样保留
        else:
            indent = int(round((l["x0"] - base_x) / char_w))
            raw_lines.append(" " * max(indent, 0) + t.rstrip())
    while raw_lines and not raw_lines[0]:
        raw_lines.pop(0)
    while raw_lines and not raw_lines[-1]:
        raw_lines.pop()
    body: List[str] = []
    blank = False
    for r in raw_lines:
        if not r:
            if blank:
                continue  # 收敛连续空行
            blank = True
        else:
            blank = False
        body.append(r)
    return "```\n" + "\n".join(body) + "\n```"


def _pdf_normalize_bullet(text: str) -> Tuple[str, bool]:
    """行首 bullet 字符归一化（痛点 4 文本侧）：强符号直接识别，弱符号需后随空格。"""
    s = text.lstrip()
    if not s:
        return text, False
    if s[0] in _PDF_BULLET_STRONG:
        return s[1:].lstrip(), True
    if s[0] in _PDF_BULLET_WEAK and len(s) > 1 and s[1] in (" ", "\u3000", "\t"):
        return s[2:].lstrip(), True
    return text, False


def _pdf_join_lines(lines: List[dict]) -> str:
    """普通段落行重组：CJK 直接拼接、西文按空格拼接、行尾连字符去除断词。"""
    buf = ""
    for l in lines:
        t = l["text"].strip()
        if not t:
            continue
        if not buf:
            buf = t
        elif buf.endswith("-") and t[:1].isascii() and t[:1].isalpha() and t[:1].islower():
            buf = buf[:-1] + t
        elif _PDF_CJK_RE.search(buf[-1:]) and _PDF_CJK_RE.search(t[:1]):
            buf += t
        else:
            buf += " " + t
    return re.sub(r"[ \t]{2,}", " ", buf)


def _pdf_point_in_rects(x: float, y: float,
                        rects: List[Tuple[float, float, float, float]]) -> bool:
    return any(x0 <= x <= x1 and y0 <= y <= y1 for (x0, y0, x1, y1) in rects)


def _pdf_render_page(pdict: dict, page_w: float, page_h: float, drawings: List[dict],
                     drop_zone: set, drop_rot: set) -> str:
    """单页重建：代码块容器 → 段落归并 → 列表符号还原 → 提示框标记，按阅读序输出。"""
    bg_rects, stroke_rects, dots = _pdf_analyze_drawings(drawings, page_w, page_h)
    blocks = _pdf_collect_blocks(pdict, page_h, drop_zone, drop_rot)
    if not blocks:
        return ""

    units: List[Tuple[int, str]] = []  # (首块索引 = 阅读序键, 输出文本)
    consumed: set = set()

    # --- 痛点 3：描边容器内的成组代码块 ---
    rect_groups: Dict[int, List[int]] = {}
    for i, b in enumerate(blocks):
        cx, cy = (b["x0"] + b["x1"]) / 2.0, (b["y0"] + b["y1"]) / 2.0
        best, best_area = -1, None
        for ri, (rx0, ry0, rx1, ry1) in enumerate(stroke_rects):
            if rx0 <= cx <= rx1 and ry0 <= cy <= ry1:
                a = (rx1 - rx0) * (ry1 - ry0)
                if best_area is None or a < best_area:
                    best, best_area = ri, a
        if best >= 0:
            rect_groups.setdefault(best, []).append(i)
    for idxs in rect_groups.values():
        group = [blocks[i] for i in idxs]
        if _pdf_group_is_code(group):
            units.append((min(idxs), _pdf_render_code(group)))
            consumed.update(idxs)
    # 等宽特征明显但无容器的独立代码块（如纯文本 diff 片段）
    for i, b in enumerate(blocks):
        if i in consumed or not b["total_chars"]:
            continue
        if b["mono_chars"] / b["total_chars"] >= 0.5 and b["total_chars"] >= 4:
            units.append((i, _pdf_render_code([b])))
            consumed.add(i)

    # --- 普通段落：相邻同缩进块合并（跨块断行重组）---
    paras: List[dict] = []
    for i, b in enumerate(blocks):
        if i in consumed:
            continue
        if paras:
            last = paras[-1]
            gap = b["y0"] - last["y1"]
            if (abs(b["x0"] - last["x0"]) <= 3.0
                    and -1.0 <= gap <= 0.55 * max(last["size"], b["size"]) + 2.0):
                last["lines"].extend(b["lines"])
                last["y1"] = b["y1"]
                last["size"] = max(last["size"], b["size"])
                continue
        paras.append({"x0": b["x0"], "y0": b["y0"], "x1": b["x1"], "y1": b["y1"],
                      "size": b["size"], "lines": list(b["lines"]), "idx": i})

    used_dots: set = set()
    for p in paras:
        text = _pdf_join_lines(p["lines"])
        if not text.strip():
            continue
        text, is_bullet = _pdf_normalize_bullet(text)
        # 痛点 4：矢量圆点与行首坐标对齐 → 还原为 "- "
        body = max(p["size"], 8.0)
        for di, (dx, dy) in enumerate(dots):
            if di in used_dots:
                continue
            if (p["y0"] - body * 0.4 <= dy <= p["y1"] + body * 0.4
                    and dx < p["x0"] and (p["x0"] - dx) <= 2.2 * body):
                used_dots.add(di)
                is_bullet = True
                break
        if is_bullet:
            text = "- " + text
        # 痛点 1：提示框（实心背景矩形）内的文字 → 输出为独立 blockquote 段落
        pcx, pcy = (p["x0"] + p["x1"]) / 2.0, (p["y0"] + p["y1"]) / 2.0
        if _pdf_point_in_rects(pcx, pcy, bg_rects):
            text = "> " + text
        units.append((p["idx"], text))

    units.sort(key=lambda u: u[0])
    return "\n\n".join(t for _, t in units)


def parse_pdf(path: Path) -> str:
    """
    解析 PDF：PyMuPDF dict 模式逐页提取并重建富文本结构（段落 / 上下标 /
    代码块 / 列表 / 提示框），处理策略见 _pdf_render_page 及各辅助函数注释。
    延迟导入 fitz：PyMuPDF 缺失时只影响 PDF 单文件（异常由 parse_file
    转为该文件的 error 状态），不在模块导入期拖垮 docx/xlsx 等其他格式。
    注意：仅能提取文本层；纯扫描件（页面全是图片）会得到空内容，需 OCR 才能处理。
    """
    try:
        # PyMuPDF >= 1.24 推荐 import pymupdf；旧版（1.23.x）只有 fitz 入口
        try:
            import pymupdf as fitz
        except ImportError:
            import fitz  # type: ignore
    except ImportError as e:
        raise RuntimeError(
            "PDF 解析依赖 PyMuPDF 未安装，请执行 pip install PyMuPDF"
        ) from e

    with fitz.open(str(path)) as doc:
        # 第一遍：缓存各页 dict/绘图对象，并跨页统计页眉页脚与旋转水印
        pages: List[Tuple[dict, float, float, List[dict]]] = [
            (page.get_text("dict"), page.rect.width, page.rect.height, page.get_drawings())
            for page in doc
        ]
    n_pages = len(pages)
    zone_counter: Counter = Counter()
    rot_counter: Counter = Counter()
    for pdict, _w, h, _dr in pages:
        for b in pdict.get("blocks", []):
            if b.get("type") != 0:
                continue
            for line in b.get("lines", []):
                norm = re.sub(r"\s+", "", "".join(
                    s.get("text", "") for s in line.get("spans", [])))
                if not norm:
                    continue
                bbox = line.get("bbox") or (0, 0, 0, 0)
                if bbox[3] <= h * 0.07 or bbox[1] >= h * 0.93:
                    zone_counter[norm] += 1
                if abs(line.get("dir", (1.0, 0.0))[1]) > 0.05:
                    rot_counter[norm] += 1
    # 页眉页脚：多页文档中出现于 60% 以上页面才剔除；单页文档不剔除（无法证实重复）
    zone_thresh = max(2, int(n_pages * 0.6)) if n_pages >= 2 else 10 ** 9
    drop_zone = {t for t, c in zone_counter.items() if c >= zone_thresh}
    drop_rot = {t for t, c in rot_counter.items() if c >= 2}

    parts = []
    for page_num, (pdict, w, h, drawings) in enumerate(pages, 1):
        text = _pdf_render_page(pdict, w, h, drawings, drop_zone, drop_rot)
        if text.strip():
            parts.append(f"[第{page_num}页]\n{text.strip()}")
    return "\n\n".join(parts).strip()


def parse_docx(path: Path) -> str:
    """解析 DOCX：段落 + 普通表格。python-docx 打开失败时降级为原始 XML 提取。"""
    try:
        doc = Document(str(path))
    except Exception as e:
        logging.warning("python-docx 打开失败 %s: %s，尝试原始 XML 兜底解析", path, e)
        return _parse_docx_raw(path)
    parts = []
    for p in doc.paragraphs:
        text = p.text.strip()
        if text:
            parts.append(text)
    for idx, table in enumerate(doc.tables, 1):
        parts.append(f"[表格{idx}]")
        for row in table.rows:
            cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))
    return "\n".join(parts).strip()


def parse_xlsx(path: Path) -> str:
    """
    解析 XLSX：每个 sheet 的单元格文本。
    注意：不使用 read_only 模式。部分第三方系统导出的 xlsx 内部 dimension
    标签谎报范围（如 2 万行数据却声明只有 A1），read_only 模式会信任该标签
    导致内容被截断甚至读空；普通模式按实际单元格扫描，可正确读出全量数据。
    """
    wb = load_workbook(filename=str(path), read_only=False, data_only=True)
    parts = []
    try:
        for ws in wb.worksheets:
            parts.append(f"[工作表] {ws.title}")
            for row in ws.iter_rows(values_only=True):
                vals = ["" if v is None else str(v) for v in row]
                if any(v.strip() for v in vals):
                    parts.append(" | ".join(vals).rstrip())
            parts.append("")
    finally:
        wb.close()
    result = "\n".join(parts).strip()
    if not result:
        logging.warning("openpyxl 未能从 %s 读到任何工作表（可能索引损坏），尝试原始 XML 兜底解析", path)
        return _parse_xlsx_raw(path)
    return result


def _is_tag(tag: str, local: str) -> bool:
    return tag == local or tag.endswith("}" + local)


def _col_index(ref: str) -> int:
    idx = 0
    for ch in ref:
        if ch.isalpha():
            idx = idx * 26 + (ord(ch.upper()) - ord("A") + 1)
        else:
            break
    return idx


def _cell_text(c: Any, shared: List[str]) -> str:
    ctype = c.get("t")
    if ctype == "inlineStr":
        return "".join(t.text or "" for t in c.iter() if _is_tag(t.tag, "t"))
    v = next((child for child in c if _is_tag(child.tag, "v")), None)
    text = (v.text or "") if v is not None else ""
    if ctype == "s" and text:
        try:
            text = shared[int(text)]
        except (ValueError, IndexError):
            text = ""
    return text


def _parse_xlsx_raw(path: Path) -> str:
    """
    兜底解析：openpyxl 因索引损坏或结构非标读不到工作表时，直接解包读原始 XML。
    不依赖命名空间与行标签：全文件抓取 <c> 单元格标签，按每个单元格自带的
    坐标（如 A12）反推行列重组表格，兼容命名空间缺失、行标签缺失、
    单元格乱序等非标准导出。
    """
    with zipfile.ZipFile(path) as z:
        names = z.namelist()
        shared: List[str] = []
        if "xl/sharedStrings.xml" in names:
            ss_root = ET.fromstring(z.read("xl/sharedStrings.xml"))
            for el in ss_root.iter():
                if _is_tag(el.tag, "si"):
                    shared.append(
                        "".join(t.text or "" for t in el.iter() if _is_tag(t.tag, "t"))
                    )
        sheet_names = sorted(n for n in names if re.match(r"xl/worksheets/sheet\d+\.xml", n))
        parts = []
        has_cell = False
        for sn in sheet_names:
            parts.append(f"[工作表] {Path(sn).stem}")
            root = ET.fromstring(z.read(sn))
            rows: Dict[int, List[Tuple[int, str]]] = {}
            for el in root.iter():
                if not _is_tag(el.tag, "c"):
                    continue
                m = re.match(r"([A-Za-z]+)(\d+)", el.get("r") or "")
                if not m:
                    continue
                text = _cell_text(el, shared).strip()
                if not text:
                    continue
                rows.setdefault(int(m.group(2)), []).append((_col_index(m.group(1)), text))
            for row_no in sorted(rows):
                has_cell = True
                parts.append(" | ".join(text for _, text in sorted(rows[row_no])))
            parts.append("")
        if not has_cell:
            raise ValueError(
                "文件的工作表主体为空（导出文件不完整：数据字典存在但单元格缺失），无法解析出内容"
            )
    return "\n".join(parts).strip()


def parse_file(path: Path, ext: str) -> Tuple[str, str, Optional[str]]:
    """
    统一解析入口。
    返回: status, content, error
    status: success / unsupported / empty / error
    """
    try:
        if ext == ".docx":
            content = parse_docx(path)
        elif ext == ".xlsx":
            content = parse_xlsx(path)
        elif ext in (".txt", ".md"):
            content = read_text_file(path)
        elif ext == ".pdf":
            content = parse_pdf(path)
        else:
            return "unsupported", "", f"不支持的文件格式：{ext}"
        if not content or not content.strip():
            return "empty", "", "文件解析后内容为空"
        return "success", content, None
    except Exception as e:
        logging.exception("解析文件异常: %s", path)
        return "error", "", f"解析失败：{type(e).__name__}: {e}"


async def process_upload_files(
    target_files: List[UploadFile],
    template_files: List[UploadFile],
) -> List[Dict[str, Any]]:
    """
    保存并解析 target / template 两类文件，返回每个文件的解析状态与文本内容。
    模板文件同样保存到 TEMPLATE_DIR 并登记进 UPLOAD_META，供导出时按模板名取用。
    返回结构遵守 UploadResponseItem 契约：filename / status / content / role / error。
    """
    results: List[Dict[str, Any]] = []
    for role, files in (("target", target_files), ("template", template_files)):
        for uf in files or []:
            original = safe_filename(uf.filename or "unnamed")
            ext = Path(original).suffix.lower()

            # 格式不支持，不保存
            if ext not in ALLOWED_EXT:
                results.append({
                    "filename": original,
                    "status": "unsupported",
                    "content": "",
                    "role": role,
                    "error": f"不支持的文件格式：{ext or '未知'}，仅支持 DOCX/TXT/MD/XLSX/PDF",
                })
                continue

            error_msg: Optional[str] = None
            asset_id: Optional[int] = None
            try:
                # 阻塞的磁盘 IO / 解析放入线程池，避免卡住事件循环
                file_id, save_path, original, ext = await run_in_threadpool(
                    save_upload_file, uf, role
                )
                status, content, error_msg = await run_in_threadpool(parse_file, save_path, ext)
                asset_id = (UPLOAD_META.get(file_id) or {}).get("asset_id")
            except Exception as e:
                logging.exception("文件上传保存失败: %s", original)
                status, content = "error", ""
                error_msg = f"上传保存失败：{type(e).__name__}: {e}"

            results.append({
                "filename": original,
                "status": status,
                "content": content if status == "success" else "",
                "role": role,
                "error": error_msg or "",
                "asset_id": asset_id,  # 增量字段：文件资产库 id，供推荐/方案引用
            })
    return results
