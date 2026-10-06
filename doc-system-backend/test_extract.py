import requests
import json
import re
import openpyxl

# 1. 文件路径
docx_path = r"C:\Users\wn153\Desktop\2025年中国城市经济百强全景报告.docx"
xlsx_path = r"C:\Users\wn153\Desktop\2025年中国城市经济百强全景报告-模板.xlsx"
output_path = r"C:\Users\wn153\Desktop\2025年中国城市经济百强全景报告-提取结果.xlsx"

# 2. 上传文件获取内容
upload_url = "http://127.0.0.1:8000/api/upload"
print("正在上传文件...")
with open(docx_path, "rb") as f1, open(xlsx_path, "rb") as f2:
    files = {
        "target_files": ("2025年中国城市经济百强全景报告.docx", f1, "application/vnd.openxmlformats-officedocument.wordprocessingml.document"),
        "template_files": ("2025年中国城市经济百强全景报告-模板.xlsx", f2, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
    }
    resp = requests.post(upload_url, files=files, timeout=30)
    resp.raise_for_status()
    upload_result = resp.json()

print("上传成功！")

# 3. 提取文档内容
content_text = ""
for item in upload_result:
    if item["role"] == "target":
        content_text = item["content"]
        break

print(f"文档内容长度: {len(content_text)} 字符")

# 4. 正则表达式提取城市数据
pattern = re.compile(
    r'(\S+) GDP 总量 ([\d,]+\.?\d*)\s*亿元（修正值）?\s*，'
    r'常住人口 ([\d,]+\.?\d*) 万，'
    r'人均 GDP ([\d,]+\.?\d*) 元，'
    r'一般公共预算收入 ([\d,]+\.?\d*) 亿元'
)

matches = pattern.findall(content_text)
print(f"\n成功匹配到 {len(matches)} 个城市数据")

# 5. 加载模板并写入数据
wb = openpyxl.load_workbook(xlsx_path)
ws = wb.active

for i, match in enumerate(matches):
    city, gdp, pop, per_gdp, revenue = match
    row = i + 2  # 第2行开始（第1行是表头）
    ws.cell(row=row, column=1, value=city)
    ws.cell(row=row, column=2, value=float(gdp.replace(",", "")))
    ws.cell(row=row, column=3, value=float(pop.replace(",", "")))
    ws.cell(row=row, column=4, value=float(per_gdp.replace(",", "")))
    ws.cell(row=row, column=5, value=float(revenue.replace(",", "")))

# 6. 保存结果
wb.save(output_path)
print(f"\n提取完成！共 {len(matches)} 个城市")
print(f"结果已保存到: {output_path}")

# 打印前5条验证
print("\n=== 前5条数据预览 ===")
for i, match in enumerate(matches[:5]):
    city, gdp, pop, per_gdp, revenue = match
    print(f"  {i+1}. {city} | GDP: {gdp} | 人口: {pop} | 人均GDP: {per_gdp} | 预算收入: {revenue}")

print("\n=== 最后3条数据预览 ===")
for i, match in enumerate(matches[-3:]):
    city, gdp, pop, per_gdp, revenue = match
    print(f"  {len(matches)-2+i}. {city} | GDP: {gdp} | 人口: {pop} | 人均GDP: {per_gdp} | 预算收入: {revenue}")