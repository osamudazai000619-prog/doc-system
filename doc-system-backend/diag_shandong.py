import sys
from pathlib import Path

sys.stdout.reconfigure(encoding="utf-8")

from dotenv import load_dotenv
load_dotenv()

from openpyxl import load_workbook
from services.upload_service import parse_xlsx
from services.extract_service import extract_date_range, filter_content_by_date

BASE = Path(r"e:\QQ\赛题二测试集\包含模板文件\2025山东省环境空气质量监测数据信息")
XLSX = BASE / "山东省环境空气质量监测数据信息202512171921_0.xlsx"
REQ = BASE / "用户要求.txt"

prompt = REQ.read_text(encoding="utf-8")

content = parse_xlsx(XLSX)
lines = content.split("\n")
print(f"[1] 后端解析结果总行数: {len(lines)}")

sheets, cur, stats = [], None, {}
for ln in lines:
    if ln.startswith("[工作表]"):
        cur = ln
        sheets.append(cur)
        stats[cur] = {"rows": 0, "header": None}
    elif cur and " | " in ln and ln.strip():
        stats[cur]["rows"] += 1
        if stats[cur]["header"] is None:
            stats[cur]["header"] = ln
print(f"[2] 工作表数: {len(sheets)}")
for s in sheets[:80]:
    st = stats[s]
    print(f"    {s}  数据行={st['rows']}")
    print(f"        表头: {st['header']}")

dr = extract_date_range(prompt)
print(f"[3] 提示词识别到的日期范围: {dr}")
if dr:
    after = filter_content_by_date(content, dr[0], dr[1])
    print(f"    日期筛选后总行数: {len(after.split(chr(10)))}")
else:
    print("    未识别到日期，日期筛选未生效！")

wb = load_workbook(str(XLSX), read_only=False, data_only=True)
semi, nl, samples = {}, 0, []
for ws in wb.worksheets:
    header = None
    for row in ws.iter_rows(values_only=True):
        vals = ["" if v is None else str(v) for v in row]
        if not any(v.strip() for v in vals):
            continue
        if header is None:
            header = vals
            continue
        for ci, v in enumerate(vals):
            if "\n" in v or "\r" in v:
                nl += 1
            if ";" in v or "；" in v:
                colname = header[ci] if ci < len(header) else f"列{ci}"
                semi[(ws.title, colname)] = semi.get((ws.title, colname), 0) + 1
        text = " | ".join(vals)
        if any(k in text for k in ("潍坊", "德州", "临沂")) and len(samples) < 6:
            samples.append(f"[{ws.title}] {text[:150]}")
wb.close()

print(f"[4] 含分号的单元格（按 工作表/列 统计）: 共{sum(semi.values())}个")
for (sheet, col), n in sorted(semi.items()):
    print(f"    {sheet} / {col}: {n}")
print(f"[5] 含换行符的单元格: {nl} 个")
print("[6] 三城市数据行样例:")
for s in samples:
    print(f"    {s}")